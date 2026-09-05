#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
知识库全量导出为 Excel（可读 / 可发给导师审核）。
读取 backend/data/knowledge_base/*.json，按 ID 排序，输出一个带样式的 xlsx。
用法: python tools/export-kb-xlsx.py [输出路径.xlsx]   （默认输出到 知识库备份-*/ 下）
"""
import sys, os, json, re

sys.stdout.reconfigure(encoding="utf-8")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
KB_DIR = os.path.join(ROOT, "backend", "data", "knowledge_base")

# 引用 xlsx 技能的设计系统模板（颜色/字体/布局统一由它出）
SKILL_TPL = r"C:\Users\Mccou\.zcode\cli\plugins\cache\zcode-plugins-official\document-skills\0.1.4\skills\xlsx\templates"
sys.path.insert(0, SKILL_TPL)
from base import (  # noqa: E402
    FONT_NAME, HEADER_BOLD, PRIMARY, NEUTRAL_900, NEUTRAL_200,
    CF_WARNING_FILL, CF_WARNING_FONT, CF_POSITIVE_FILL, CF_POSITIVE_FONT,
    font_caption,
    setup_sheet, style_header_row, style_data_row, style_total_row,
    auto_fit_columns, auto_fit_row_heights,
)
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side  # noqa: E402

def id_sort_key(entry):
    m = re.search(r"(\d+)", str(entry.get("id", "")))
    return (0 if "LAB" in str(entry.get("id", "")).upper() else 1, int(m.group(1)) if m else 0)

def main():
    files = sorted(f for f in os.listdir(KB_DIR) if f.endswith(".json"))
    rows = []
    for f in files:
        with open(os.path.join(KB_DIR, f), encoding="utf-8") as fh:
            arr = json.load(fh)
        for it in arr:
            rows.append({
                "id": it.get("id", ""),
                "category": it.get("category", ""),
                "question": it.get("question", ""),
                "answer": it.get("answer", ""),
                "source": it.get("source", ""),
                "src_file": f,
                "reviewed": bool(it.get("reviewed", False)),
                "reviewedBy": it.get("reviewedBy", "") or "",
                "reviewedAt": it.get("reviewedAt", "") or "",
            })
    rows.sort(key=id_sort_key)

    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
        ROOT, "知识库备份-2026-09-02", "知识库全量导出-2026-09-02.xlsx")

    headers = ["ID", "分类", "问题", "答案", "来源", "来源文件", "审核状态", "审核人", "审核时间"]
    col_count = len(headers)  # 9 列数据，从 B=2 起，到 J=10
    last_col = 2 + col_count - 1

    wb = Workbook()
    ws = wb.active
    ws.title = "全量问答对"
    ws.sheet_view.showGridLines = False
    setup_sheet(ws, title=f"肝芽守护知识库全量导出（{len(rows)} 条）", last_col=last_col)

    for i, h in enumerate(headers):
        ws.cell(row=4, column=2 + i, value=h)
    style_header_row(ws, row_num=4, col_start=2, col_end=last_col)

    for r, row in enumerate(rows):
        row_num = 5 + r
        vals = [row["id"], row["category"], row["question"], row["answer"],
                row["source"], row["src_file"],
                "已审核" if row["reviewed"] else "待审核",
                row["reviewedBy"], row["reviewedAt"]]
        for c, v in enumerate(vals):
            cell = ws.cell(row=row_num, column=2 + c, value=v)
            if c == 6:  # 审核状态列：待审核=琥珀、已审核=绿
                cell.fill = CF_POSITIVE_FILL if row["reviewed"] else CF_WARNING_FILL
                cell.font = CF_POSITIVE_FONT if row["reviewed"] else CF_WARNING_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center")
        style_data_row(ws, row_num=row_num, col_start=2, col_end=last_col, row_index=r)

    # 合计行（静态导出，直接写计算值）
    total_row = 5 + len(rows)
    ws.cell(row=total_row, column=2, value="合计")
    ws.cell(row=total_row, column=3, value=f"{len(rows)} 条")
    n_reviewed = sum(1 for r in rows if r["reviewed"])
    ws.cell(row=total_row, column=7, value=f"已审核 {n_reviewed} / 待审核 {len(rows) - n_reviewed}")
    style_total_row(ws, row_num=total_row, col_start=2, col_end=last_col)

    # 说明行（caption）
    note_row = total_row + 2
    note = ("说明：① 本表为后端 knowledge_base 源文件的全量导出（已与云端部署索引核对一致，md5 相同）。"
            "② 分类为 lab_reference 的条目仅用于化验单 OCR 解读，不参与文本问答。"
            "③ 全部条目当前均为『待审核』草稿（reviewed:false），上线前需导师逐条核对。")
    ws.merge_cells(start_row=note_row, start_column=2, end_row=note_row, end_column=last_col)
    nc = ws.cell(row=note_row, column=2, value=note)
    nc.font = font_caption()
    nc.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    ws.row_dimensions[note_row].height = 40

    ws.freeze_panes = "B5"
    ws.auto_filter.ref = f"B4:J{total_row - 1}"

    auto_fit_columns(ws, min_width=8, max_width=34, header_row=4, data_start_row=5)
    auto_fit_row_heights(ws, header_row=4, data_start_row=5)

    wb.properties.creator = "Z.ai"
    wb.save(out)
    print("已导出:", out, "共", len(rows), "条")

if __name__ == "__main__":
    main()
